import io
import json
import sys
import unittest
import urllib.error
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import engine
from docx import Document

class EngineTests(unittest.TestCase):
    def setUp(self):
        self.docs,self.rates=engine.load_knowledge([])
        self.price=engine.pricing(self.rates,{'discovery':2,'training':1},10,18,'INR')
        self.request={'client':'Example Client','industry':'Manufacturing',
                      'requirements':'Track CRM pipeline leads and train users.',
                      'timeline':'6 weeks','services':['Discovery','Training']}
        self.settings={'demo':True,'model_id':'Demo','temperature':0.7,'token_cap':600}
        self.bundle=engine.run_engineered(self.request,self.price,[],self.settings,None)

    def test_hand_checked_price(self):
        self.assertEqual(self.price['grand_total'],'26550.00')

    def test_invalid_price_inputs(self):
        for quantities,currency in [({'training':-1},'INR'),({'training':1},'USD')]:
            with self.assertRaises(ValueError):
                engine.pricing(self.rates,quantities,0,0,currency)

    def test_retrieval_without_models(self):
        retriever=engine.Retriever({'crm.txt':'CRM pipeline leads','plants.txt':'flowers petals leaves'},False)
        self.assertEqual(retriever.search('CRM pipeline')[0]['source'],'crm.txt')

    def test_currency_mismatch_detection(self):
        sections=dict(self.bundle['sections'])
        sections['Pricing Assumptions']+=' INR 26,550.00'
        self.assertFalse(engine.evaluate(sections,self.request,self.price)['mismatches'])
        sections['Email Draft']+=' INR 999.00 USD 26,550.00'
        self.assertEqual(len(engine.evaluate(sections,self.request,self.price)['mismatches']),2)

    def test_export_and_baseline(self):
        self.assertTrue(all(self.bundle['sections'].get(n) for n in engine.SECTIONS))
        document=Document(io.BytesIO(engine.export_docx(self.bundle)))
        self.assertEqual(len(document.tables),1)
        self.assertIn('Grand total',engine.export_markdown(self.bundle))
        baseline=engine.run_baseline(self.request,self.price,[],self.settings,None)
        self.assertTrue(baseline['evaluation']['placeholders'])

    def test_text_supplement_and_full_replacement(self):
        docs,_=engine.load_knowledge([('custom.txt',b'Authorised company facts')])
        self.assertIn('UPLOAD/custom.txt',docs)
        with self.assertRaises(ValueError):
            engine.load_knowledge([('custom.txt',b'Company facts')],True)
        docs,rates=engine.load_knowledge([('custom.txt',b'Company facts'),('services_and_rates.json',json.dumps(self.rates).encode())],True)
        self.assertFalse(any(k.startswith('FICTIONAL') for k in docs))
        self.assertEqual(rates['currency'],'INR')

    def test_hosted_adapter_contract(self):
        settings=dict(self.settings,demo=False,model_id='test/model')
        body=json.dumps({'choices':[{'finish_reason':'stop','message':{'content':'Draft text'}}]}).encode()
        with patch('urllib.request.urlopen',return_value=io.BytesIO(body)) as call:
            output=engine.ModelManager('test-token').generate([{'role':'user','content':'Test'}],settings,300)
        self.assertEqual(output,'Draft text')
        request=call.call_args.args[0]
        payload=json.loads(request.data)
        self.assertEqual(request.full_url,engine.ModelManager.ENDPOINT)
        self.assertEqual(payload['max_tokens'],300)
        self.assertNotIn('test-token',json.dumps(payload))
        self.assertEqual(payload['model'],'test/model')

    def test_hosted_error_redacts_provider_body(self):
        error=urllib.error.HTTPError(engine.ModelManager.ENDPOINT,401,'Secret provider details',{},None)
        with patch('urllib.request.urlopen',side_effect=error):
            with self.assertRaisesRegex(RuntimeError,'rejected the token') as caught:
                engine.ModelManager('test-token').generate([],dict(self.settings,model_id='test/model'),250)
        self.assertNotIn('Secret provider details',str(caught.exception))

    def test_no_request_without_token(self):
        with patch('urllib.request.urlopen') as call:
            with self.assertRaises(ValueError):
                engine.ModelManager('').generate([],self.settings,250)
            call.assert_not_called()

if __name__=='__main__':
    unittest.main()
